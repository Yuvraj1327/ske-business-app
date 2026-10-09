"""Builds the Sales Return .xlsx export."""
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = ("Customer Name", "Invoice Number", "Total Amount")


def _text(ws, row: int, column: int, value: str) -> None:
    # Force a literal string cell: openpyxl would otherwise store values that
    # start with "=" as formulas (spreadsheet formula injection via a name).
    cell = ws.cell(row=row, column=column, value=value)
    cell.data_type = "s"


def build_sales_returns_xlsx(rows: list[tuple[str, str, Decimal]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Returns"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, (customer_name, invoice_number, amount) in enumerate(rows, start=2):
        _text(ws, i, 1, customer_name)
        _text(ws, i, 2, invoice_number)
        amount_cell = ws.cell(row=i, column=3, value=amount)
        amount_cell.number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
