"""Builds the .xlsx downloads (Sales Return export, Settlement cash/online export)."""
import io
import re
from collections.abc import Sequence
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _text(ws, row: int, column: int, value: str) -> None:
    # Force a literal string cell: openpyxl would otherwise store values that
    # start with "=" as formulas (spreadsheet formula injection via a name).
    cell = ws.cell(row=row, column=column, value=value)
    cell.data_type = "s"


def build_xlsx(
    sheet_title: str, headers: Sequence[str], rows: Sequence[tuple[str, str, Decimal]], widths: Sequence[int] = (32, 22, 16)
) -> bytes:
    """One header row + `rows` of (text, text, amount) — i.e. Customer Name,
    Invoice Number, <some> Amount."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(list(headers))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, (first, second, amount) in enumerate(rows, start=2):
        _text(ws, i, 1, first)
        _text(ws, i, 2, second)
        amount_cell = ws.cell(row=i, column=3, value=amount)
        amount_cell.number_format = "#,##0.00"
    for letter, width in zip("ABC", widths):
        ws.column_dimensions[letter].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def safe_filename_part(value: str) -> str:
    """Keeps a user-entered value (e.g. a Pick Sheet No.) safe inside a filename."""
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._") or "sheet"
