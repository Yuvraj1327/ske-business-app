"""Settlement sheet Online / Cash .xlsx export against a REAL (temporary
SQLite file) database: a picklist imported through the real import code,
payments saved through the real PicklistService.confirm_item, the sheet made
through the real SettlementService, and the real route (see
test_picklist_settlement_sync_db.py for the shared setup)."""
import io
import uuid
from datetime import date
from decimal import Decimal

import httpx
import pytest
from openpyxl import load_workbook

from app.core.security import CurrentUser, get_current_user
from app.db.session import get_db
from app.main import app
from app.schemas.settlement import SettlementSheetCreateRequest
from app.services.settlement_service import SettlementService
from app.tests.test_picklist_settlement_sync_db import (  # noqa: F401  (session_factory is a fixture)
    PICKLIST_NO,
    _confirm,
    _item_ids,
    _seed,
    _user,
    session_factory,
)

D = Decimal
URL = "/api/v1/settlements/{}/export"


def _client(factory, user: CurrentUser):
    async def _db():
        async with factory() as db:
            yield db

    async def _me():
        return user

    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_current_user] = _me
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t")


@pytest.fixture(autouse=True)
def _clear_overrides():
    yield
    app.dependency_overrides.clear()


def _rows(content: bytes):
    ws = load_workbook(io.BytesIO(content)).active
    return [tuple(c.value for c in r) for r in ws.iter_rows()]


async def _sheet(factory, ids, admin, pick_sheet_no=PICKLIST_NO) -> uuid.UUID:
    async with factory() as db:
        created = await SettlementService(db).create_sheet(
            SettlementSheetCreateRequest(
                sheet_date=date(2026, 10, 9), delivery_agent_id=ids["delivery_agent"],
                salesman_ids=[ids["salesman"]], pick_sheet_no=pick_sheet_no,
            ),
            admin,
        )
    return created.id


@pytest.mark.asyncio
async def test_online_and_cash_exports_match_saved_payments(session_factory):
    f = session_factory
    ids = await _seed(f)
    admin = _user(ids["admin"], "admin", "settlements.manage")
    agent = _user(ids["delivery_agent"], "delivery_agent", "picklists.view_assigned", "settlements.view_assigned")
    sheet_id = await _sheet(f, ids, admin)

    items = await _item_ids(f)
    await _confirm(f, items[0], "cash", agent)      # Customer 1 / INV-1 / 1000
    await _confirm(f, items[1], "online", agent)    # Customer 2 / INV-2 / 2000
    await _confirm(f, items[2], "credit", agent, salesman_id=ids["salesman"])
    await _confirm(f, items[3], "cheque", agent, cheque_amount=D("300"))
    await _confirm(f, items[4], "cash", agent)      # Customer 5 / INV-5 / 300

    async with _client(f, admin) as c:
        online = await c.get(URL.format(sheet_id), params={"mode": "online"})
        cash = await c.get(URL.format(sheet_id), params={"mode": "cash"})
        sheet = (await c.get(f"/api/v1/settlements/{sheet_id}")).json()

    for resp, mode in ((online, "online"), (cash, "cash")):
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert f"{PICKLIST_NO}_{mode}_{date.today().isoformat()}.xlsx" in resp.headers["content-disposition"]

    on = _rows(online.content)
    assert on == [("Customer Name", "Invoice Number", "Online Amount"), ("Customer 2", "INV-2", 2000)]
    ca = _rows(cash.content)
    assert ca == [
        ("Customer Name", "Invoice Number", "Cash Amount"),
        ("Customer 1", "INV-1", 1000),
        ("Customer 5", "INV-5", 300),
    ]
    # The exported amounts add up to the totals the Settlement sheet itself shows.
    assert D(str(sum(r[2] for r in on[1:]))) == D(sheet["online_amount"])
    assert D(str(sum(r[2] for r in ca[1:]))) == D(sheet["cash_amount"])


@pytest.mark.asyncio
async def test_export_only_includes_this_sheets_picklist_and_reports_empty(session_factory):
    f = session_factory
    ids = await _seed(f)
    admin = _user(ids["admin"], "admin", "settlements.manage")
    sheet_id = await _sheet(f, ids, admin)
    no_link = await _sheet(f, ids, admin, pick_sheet_no=None)
    other = await _sheet(f, ids, admin, pick_sheet_no="PL-DOES-NOT-EXIST")

    async with _client(f, admin) as c:
        # Picklist exists but nothing saved as online yet -> clear message, no file.
        r = await c.get(URL.format(sheet_id), params={"mode": "online"})
        assert r.status_code == 400 and "No Online payments" in r.json()["error"]["message"]
        # Sheet without a Pick Sheet No. / with an unknown one -> clear errors, nothing exported.
        r = await c.get(URL.format(no_link), params={"mode": "cash"})
        assert r.status_code == 400 and "Pick Sheet No." in r.json()["error"]["message"]
        r = await c.get(URL.format(other), params={"mode": "cash"})
        assert r.status_code == 404
        # Unknown sheet / invalid mode.
        assert (await c.get(URL.format(uuid.uuid4()), params={"mode": "cash"})).status_code == 404
        assert (await c.get(URL.format(sheet_id), params={"mode": "cheque"})).status_code == 422


@pytest.mark.asyncio
async def test_export_follows_sheet_view_access(session_factory):
    f = session_factory
    ids = await _seed(f)
    admin = _user(ids["admin"], "admin", "settlements.manage")
    agent = _user(ids["delivery_agent"], "delivery_agent", "picklists.view_assigned", "settlements.view_assigned")
    sheet_id = await _sheet(f, ids, admin)
    await _confirm(f, (await _item_ids(f))[0], "cash", agent)

    # The sheet's own delivery agent can export; an unrelated user and a user with no settlement permission cannot.
    async with _client(f, agent) as c:
        assert (await c.get(URL.format(sheet_id), params={"mode": "cash"})).status_code == 200
    stranger = _user(uuid.uuid4(), "salesman", "settlements.view_assigned")
    async with _client(f, stranger) as c:
        assert (await c.get(URL.format(sheet_id), params={"mode": "cash"})).status_code == 403
    nobody = _user(ids["delivery_agent"], "delivery_agent")
    async with _client(f, nobody) as c:
        assert (await c.get(URL.format(sheet_id), params={"mode": "cash"})).status_code == 403
